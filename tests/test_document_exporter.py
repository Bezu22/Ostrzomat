import tempfile
import unittest
from pathlib import Path

from docx import Document

from utils.document_exporter import generate_docx, generate_pdf


class TestDocumentExporter(unittest.TestCase):
    def setUp(self):
        self.cart_data = {
            "items": [{
                "type": "Frez testowy",
                "diam": "10",
                "shank_diam": "10",
                "coat_len": "20",
                "z": "4",
                "qty": "2",
                "tool_unit": 5.0,
                "total_tool": 10.0,
                "coat_name": "TiN",
                "coat_unit": 2.5,
                "total_coat": 5.0,
                "services_status": {"ciecie": True},
                "total_extra": 3.0,
                "notes": "uwaga",
            }]
        }
        self.client_info = {
            "id": 1,
            "name": "Klient testowy",
            "phone": "-",
            "nip": "-",
            "email": "-",
            "address": "-",
        }

    def test_docx_contains_z_and_position_total_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "raport.docx"
            generate_docx(self.cart_data, self.client_info, str(output_path))
            document = Document(output_path)

        table = document.tables[2]
        headers = [cell.text for cell in table.rows[0].cells]
        values = [cell.text for cell in table.rows[1].cells]

        self.assertEqual(len(headers), 16)
        self.assertEqual(headers[5], "Z")
        self.assertEqual(headers[14], "Suma P")
        self.assertEqual(values[14], "18.00 zł")

    def test_docx_supports_custom_report_number_and_optional_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "raport_custom.docx"
            generate_docx(
                self.cart_data, self.client_info, str(output_path),
                report_number="115", include_client=False, include_summary=False, include_logo=False
            )
            document = Document(output_path)

        # Bez sekcji klienta i podsumowania mamy tylko tabelę nagłówka i tabelę pozycji
        self.assertEqual(len(document.tables), 2)
        header_text = document.tables[0].cell(0, 1).text
        self.assertIn("nr 115", header_text)

    def test_pdf_supports_custom_report_number_and_options(self):
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "raport_custom.pdf"
            generate_pdf(
                self.cart_data, self.client_info, str(output_path),
                report_number="115", include_client=False, include_summary=False
            )
            self.assertTrue(output_path.exists())
            self.assertGreater(output_path.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
