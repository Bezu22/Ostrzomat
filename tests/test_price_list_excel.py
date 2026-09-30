import shutil
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import database
from utils.price_list_excel import export_pricelist, import_pricelist


class TestPriceListExcel(unittest.TestCase):
    def test_excel_round_trip_supports_new_tool_and_quantity_discount(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            db_path = temp_path / "ostrzomat.db"
            excel_path = temp_path / "cennik.xlsx"
            shutil.copy2(Path("data") / "ostrzomat.db", db_path)

            original_db_path = database.DB_PATH
            database.DB_PATH = str(db_path)
            try:
                export_pricelist(excel_path)
                workbook = load_workbook(excel_path)
                workbook["Narzędzia"].append(
                    ["Frezy", "Frez testowy", 1, 4, 20.1, 25.0, 100.0]
                )
                # Zmieniamy poprzedni zakres 11+ na 11-20, aby zrobić miejsce dla 21-50
                workbook["Rabaty ilościowe"]["B5"] = 20
                workbook["Rabaty ilościowe"].append(
                    [21, 50, 30]
                )
                workbook.save(excel_path)
                workbook.close()

                counts = import_pricelist(excel_path)
                self.assertGreater(counts["Narzędzia"], 0)
                self.assertGreater(counts["Rabaty ilościowe"], 0)
                
                # Frez testowy cena bazowa 100 zł; dla partii 25 sztuk rabat 30% -> 70.0 zł
                self.assertEqual(database.get_tool_price("Frez testowy", "4", 22, 25), 70.0)
                self.assertEqual(
                    sqlite3.connect(db_path).execute(
                        "SELECT COUNT(*) FROM pricelist_tools WHERE tool_type=?", ("Frez testowy",)
                    ).fetchone()[0],
                    1,
                )
            finally:
                database.DB_PATH = original_db_path

    def test_invalid_quantity_range_reports_cells_and_preserves_cache(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            db_path = temp_path / "ostrzomat.db"
            excel_path = temp_path / "cennik.xlsx"
            shutil.copy2(Path("data") / "ostrzomat.db", db_path)

            original_db_path = database.DB_PATH
            database.DB_PATH = str(db_path)
            try:
                export_pricelist(excel_path)
                workbook = load_workbook(excel_path)
                # Ustawienie min > max w Rabaty ilościowe (np. od 30 do 20 szt.)
                workbook["Rabaty ilościowe"]["A2"] = 30
                workbook["Rabaty ilościowe"]["B2"] = 20
                workbook.save(excel_path)
                workbook.close()

                before = sqlite3.connect(db_path).execute(
                    "SELECT COUNT(*) FROM pricelist_quantity_discounts"
                ).fetchone()[0]
                with self.assertRaisesRegex(ValueError, r"Rabaty ilościowe!A2/B2"):
                    import_pricelist(excel_path)
                after = sqlite3.connect(db_path).execute(
                    "SELECT COUNT(*) FROM pricelist_quantity_discounts"
                ).fetchone()[0]
                self.assertEqual(before, after)
            finally:
                database.DB_PATH = original_db_path

    def test_editing_standard_tool_price_reflects_in_calculated_price(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            db_path = temp_path / "ostrzomat.db"
            excel_path = temp_path / "cennik.xlsx"
            shutil.copy2(Path("data") / "ostrzomat.db", db_path)

            original_db_path = database.DB_PATH
            database.DB_PATH = str(db_path)
            try:
                export_pricelist(excel_path)
                workbook = load_workbook(excel_path)
                # Kolumna G to 'Cena bazowa' w arkuszu Narzędzia
                workbook["Narzędzia"]["G2"] = 999
                workbook.save(excel_path)
                workbook.close()

                import_pricelist(excel_path)
                # Dla 1 sztuki (rabat 0%) cena wynosi dokładnie 999.0
                self.assertEqual(database.get_tool_price("Frez prosty", "1", 5, 1), 999.0)
            finally:
                database.DB_PATH = original_db_path


if __name__ == "__main__":
    unittest.main()