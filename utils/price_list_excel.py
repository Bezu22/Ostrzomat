from pathlib import Path
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

import database


DEFAULT_EXCEL_PATH = Path("data") / "cennik.xlsx"

SHEETS = {
    "Narzędzia": (
        "category", "tool_type", "blades_min", "blades_max", "diam_min", "diam_max",
        "price_1", "price_2_4", "price_5_10", "price_11_20",
    ),
    "Powłoki": ("coating_name", "diam_max", "length", "price"),
    "Usługi": ("service_name", "param_min", "param_max", "price"),
    "Zakresy ilościowe": (
        "tool_type", "blades_min", "blades_max", "diam_min", "diam_max", "qty_min", "qty_max", "price"
    ),
}

HEADERS = {
    "Narzędzia": ("Kategoria", "Typ narzędzia", "Ostrza min", "Ostrza max",
                  "Średnica min", "Średnica max", "Cena 1", "Cena 2-4", "Cena 5-10", "Cena 11+"),
    "Powłoki": ("Nazwa powłoki", "Średnica max", "Długość", "Cena"),
    "Usługi": ("Nazwa usługi", "Parametr min", "Parametr max", "Cena"),
    "Zakresy ilościowe": (
        "Typ narzędzia", "Ostrza min", "Ostrza max", "Średnica min", "Średnica max",
        "Ilość min", "Ilość max", "Cena",
    ),
}

TABLES = {
    "Narzędzia": "pricelist_tools",
    "Powłoki": "pricelist_coatings",
    "Usługi": "pricelist_services",
    "Zakresy ilościowe": "pricelist_tool_ranges",
}


def export_pricelist(path=DEFAULT_EXCEL_PATH):
    """Eksportuje aktualny cennik SQLite do jednego, edytowalnego pliku XLSX."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    connection = database.get_connection()
    try:
        database.ensure_tool_blade_columns(connection)
        connection.commit()
        for sheet_name, columns in SHEETS.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(HEADERS[sheet_name])
            table = TABLES[sheet_name]
            if sheet_name == "Zakresy ilościowe":
                database.ensure_tool_ranges_table(connection)
            if sheet_name == "Narzędzia":
                rows = connection.execute(
                    "SELECT category, tool_type, blades_min, blades_max, diam_min, diam_max, "
                    "price_1, price_2_4, price_5_10, price_11_20 "
                    "FROM pricelist_tools ORDER BY id"
                ).fetchall()
            else:
                rows = connection.execute(
                    f"SELECT {', '.join(columns)} FROM {table} ORDER BY id"
                ).fetchall()
            if sheet_name == "Zakresy ilościowe" and not rows:
                legacy_rows = connection.execute(
                    "SELECT tool_type, blades_min, blades_max, diam_min, diam_max, price_1, price_2_4, price_5_10, price_11_20 "
                    "FROM pricelist_tools ORDER BY id"
                ).fetchall()
                ranges = ((1, 1), (2, 4), (5, 10), (11, 999999))
                rows = [
                    (tool_type, blades_min, blades_max, diam_min, diam_max, qty_min, qty_max, prices[index])
                    for tool_type, blades_min, blades_max, diam_min, diam_max, *prices in legacy_rows
                    for index, (qty_min, qty_max) in enumerate(ranges)
                ]
            for row in rows:
                sheet.append(list(row))
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                font = copy(cell.font)
                font.bold = True
                cell.font = font
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = max(
                    12, min(28, max(len(str(cell.value or "")) for cell in column) + 2)
                )
    finally:
        connection.close()
    workbook.save(path)
    return path


def import_pricelist(path=DEFAULT_EXCEL_PATH):
    """Waliduje i importuje wszystkie arkusze XLSX do SQLite w jednej transakcji."""
    parsed = read_pricelist(path)

    connection = database.get_connection()
    try:
        connection.execute("BEGIN")
        for sheet_name, rows in parsed.items():
            table = TABLES[sheet_name]
            if sheet_name == "Zakresy ilościowe":
                database.ensure_tool_ranges_table(connection)
            connection.execute(f"DELETE FROM {table}")
            insert_columns = list(SHEETS[sheet_name])
            insert_rows = rows
            table_columns = {row[1] for row in connection.execute(f"PRAGMA table_info({table})")}
            if sheet_name == "Narzędzia" and "blades" in table_columns:
                insert_columns.insert(4, "blades")
                insert_rows = [row[:4] + (_blade_label(row[2], row[3]),) + row[4:] for row in rows]
            if sheet_name == "Zakresy ilościowe" and "blades" in table_columns:
                insert_columns.insert(3, "blades")
                insert_rows = [row[:3] + (_blade_label(row[1], row[2]),) + row[3:] for row in rows]
            placeholders = ", ".join("?" for _ in insert_columns)
            connection.executemany(
                f"INSERT INTO {table} ({', '.join(insert_columns)}) VALUES ({placeholders})",
                insert_rows,
            )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    return {sheet_name: len(rows) for sheet_name, rows in parsed.items()}


def read_pricelist(path=DEFAULT_EXCEL_PATH):
    """Wczytuje, typuje i sprawdza cały skoroszyt bez zapisu do bazy."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Brak pliku cennika: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        parsed = {}
        for sheet_name, columns in SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Brak arkusza '{sheet_name}'")
            rows = workbook[sheet_name].iter_rows(values_only=True)
            header = tuple(next(rows, ()))
            if header != HEADERS[sheet_name]:
                raise ValueError(f"Nieprawidłowe nagłówki arkusza '{sheet_name}'")
            parsed[sheet_name] = []
            for row_number, row in enumerate(rows, start=2):
                if any(row):
                    try:
                        if sheet_name == "Narzędzia" and len(row) == 9:
                            row = row[:2] + _normalize_blade_range((row[2], None)) + row[3:]
                        elif sheet_name == "Zakresy ilościowe" and len(row) == 7:
                            row = row[:1] + _normalize_blade_range((row[1], None)) + row[2:]
                        parsed[sheet_name].append(_parse_row(sheet_name, row, columns, row_number))
                    except ValueError as error:
                        raise ValueError(f"{sheet_name}!A{row_number}: {error}") from error
    finally:
        workbook.close()

    _validate_pricelist(parsed)
    return parsed


def _validate_pricelist(parsed):
    tool_keys = {(row[1], row[2], row[3]) for row in parsed["Narzędzia"]}
    for sheet_name, rows in parsed.items():
        for row_number, row in enumerate(rows, start=2):
            if sheet_name == "Narzędzia":
                blade_range = (2, 3)
                min_index, max_index = 4, 5
                quantity_range = None
            elif sheet_name == "Zakresy ilościowe":
                blade_range = (1, 2)
                min_index, max_index = 3, 4
                quantity_range = (5, 6)
            else:
                min_index, max_index = 1, 2
                quantity_range = None
            min_cell = get_column_letter(min_index + 1)
            max_cell = get_column_letter(max_index + 1)
            if sheet_name == "Narzędzia" and row[blade_range[0]] > row[blade_range[1]]:
                raise ValueError(f"{sheet_name}!C{row_number}/D{row_number}: Ostrza min jest większe od Ostrza max")
            if row[min_index] > row[max_index]:
                raise ValueError(f"{sheet_name}!{min_cell}{row_number}/{max_cell}{row_number}: minimum jest większe od maksimum")
            if quantity_range and row[quantity_range[0]] > row[quantity_range[1]]:
                raise ValueError(f"{sheet_name}!F{row_number}/G{row_number}: ilość min jest większa od ilości max")
            if any(value < 0 for value in row[min_index:max_index + 1]):
                raise ValueError(f"{sheet_name}!{min_cell}{row_number}:{max_cell}{row_number}: zakres nie może być ujemny")
            if quantity_range:
                if not any(
                    row[0] == tool[1] and row[1] == tool[2] and row[2] == tool[3]
                    and row[3] <= tool[5] and row[4] >= tool[4]
                    for tool in parsed["Narzędzia"]
                ):
                    raise ValueError(f"{sheet_name}!A{row_number}: brak narzędzia w arkuszu Narzędzia")
                if any(
                    value != int(value) or value < 1
                    for value in row[quantity_range[0]:quantity_range[1] + 1]
                ):
                    raise ValueError(f"{sheet_name}!E{row_number}: ilości muszą być dodatnimi liczbami całkowitymi")
            price_start = quantity_range[1] + 1 if quantity_range else max_index + 1
            if any(value < 0 for value in row[price_start:]):
                raise ValueError(f"{sheet_name}!{get_column_letter(price_start + 1)}{row_number}: cena nie może być ujemna")

    ranges = parsed["Zakresy ilościowe"]
    for index, current in enumerate(ranges):
        for other in ranges[index + 1:]:
            same_key = current[:3] == other[:3]
            diam_overlap = current[3] <= other[4] and other[3] <= current[4]
            quantity_overlap = current[5] <= other[6] and other[5] <= current[6]
            if same_key and diam_overlap and quantity_overlap:
                first_row = index + 2
                second_row = ranges.index(other) + 2
                raise ValueError(
                    f"Zakresy ilościowe!A{first_row}:G{first_row}/A{second_row}:G{second_row}: "
                    "zakresy nakładają się dla tego samego narzędzia"
                )


def _parse_row(sheet_name, row, columns, row_number):
    if len(row) != len(columns) or any(value is None for value in row):
        missing_index = next((index for index, value in enumerate(row) if value is None), 0)
        raise ValueError(f"Niepełna wartość w komórce {get_column_letter(missing_index + 1)}{row_number}")
    values = list(row)
    if sheet_name == "Narzędzia":
        values[0:2] = [str(value).strip() for value in values[0:2]]
        values[2:4] = _normalize_blade_range(values[2:4])
        numeric_start = 2
    elif sheet_name == "Zakresy ilościowe":
        values[0] = str(values[0]).strip()
        values[1:3] = _normalize_blade_range(values[1:3])
        numeric_start = 1
    else:
        values[0] = str(values[0]).strip()
        numeric_start = 1
    try:
        for index in range(numeric_start, len(values)):
            values[index] = float(values[index])
    except (TypeError, ValueError) as error:
        raise ValueError(f"Nieprawidłowa liczba w komórce {get_column_letter(index + 1)}{row_number}") from error
    return tuple(values)


def _normalize_blade_range(values):
    if len(values) == 2 and isinstance(values[0], str) and values[0] in ("2-4", "1-4", "pozostałe", "5-99"):
        return (1, 4) if values[0] in ("2-4", "1-4") else (5, 99)
    return values


def _blade_label(blades_min, blades_max):
    return f"{int(blades_min)}-{int(blades_max)}"